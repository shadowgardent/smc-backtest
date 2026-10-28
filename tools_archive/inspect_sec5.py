import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r"C:\Users\ONEST\Downloads\SMC_AUCC2026_Final.docx")

for i in range(65, 71):
    print(f"P{i:02d}: '{doc.paragraphs[i].text}'")

table3 = doc.tables[2]
print("\nTable 3 rows:")
for r in table3.rows:
    print([c.text.strip().replace("\n", " ") for c in r.cells])
