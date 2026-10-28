import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r"C:\Users\ONEST\Downloads\SMC_AUCC2026_Final.docx")
p29 = doc.paragraphs[29]
p30 = doc.paragraphs[30]

print("P29 alignment:", p29.alignment)
print("P29 space before/after:", p29.paragraph_format.space_before, p29.paragraph_format.space_after)

print("P30 text:", p30.text)
print("P30 alignment:", p30.alignment)
print("P30 space before/after:", p30.paragraph_format.space_before, p30.paragraph_format.space_after)
for r in p30.runs:
    print(f"  Run: text='{r.text}', font={r.font.name}, size={r.font.size}, bold={r.font.bold}")
