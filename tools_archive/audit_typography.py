import docx
from docx.shared import Pt, Inches
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r"C:\Users\ONEST\Downloads\SMC_AUCC2026_Final.docx")

print("=== 1. PARAGRAPH AUDIT ===")
fonts_found = set()
sizes_found = set()
line_spacings = set()
space_befores = set()
space_afters = set()
alignments = set()

anomalies = []

for i, p in enumerate(doc.paragraphs):
    # Paragraph format
    pf = p.paragraph_format
    ls = pf.line_spacing
    sb = pf.space_before.pt if pf.space_before else 0
    sa = pf.space_after.pt if pf.space_after else 0
    al = p.alignment
    
    line_spacings.add(ls)
    space_befores.add(sb)
    space_afters.add(sa)
    alignments.add(str(al))
    
    p_text = p.text.strip()
    if not p_text and not any("drawing" in r._r.xml for r in p.runs):
        anomalies.append(f"P{i:03d}: Blank paragraph!")
        
    for j, r in enumerate(p.runs):
        fn = r.font.name
        fs = r.font.size.pt if r.font.size else None
        b = r.font.bold
        it = r.font.italic
        
        fonts_found.add(fn)
        sizes_found.add(fs)
        
        # Check font consistency
        if fn and "sarabun" not in fn.lower():
            anomalies.append(f"P{i:03d} Run{j}: Non-Sarabun font '{fn}' in text '{r.text[:30]}'")
            
print("Fonts found:", fonts_found)
print("Font sizes found (pt):", sorted([s for s in sizes_found if s is not None]))
print("Line spacings found:", line_spacings)
print("Space before found (pt):", sorted(list(space_befores)))
print("Space after found (pt):", sorted(list(space_afters)))
print("Alignments found:", alignments)

print("\n=== 2. TABLE AUDIT ===")
table_fonts = set()
table_sizes = set()
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.font.name: table_fonts.add(r.font.name)
                    if r.font.size: table_sizes.add(r.font.size.pt)

print("Table fonts:", table_fonts)
print("Table sizes (pt):", sorted(list(table_sizes)))

print("\n=== 3. DETAILED BREAKDOWN BY SECTION ===")
section_checks = [
    ("Title Thai", 0),
    ("Title Eng", 1),
    ("Author", 2),
    ("Abstract Header", 5),
    ("Abstract Body", 6),
    ("Keywords", 7),
    ("Section 1 Header", 13),
    ("Section 1 Body", 14),
    ("Table 1 Caption", 20),
    ("Figure 1 Caption", 30),
    ("Section 3 Header", 23),
    ("Section 3.1 Header", 24),
    ("Figure 2 Caption", 63),
    ("Figure 3 Caption", 65),
    ("Table 3 Caption", 69),
    ("References Header", 80),
]

for label, idx in section_checks:
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        run_info = []
        for r in p.runs:
            fn = r.font.name
            fs = r.font.size.pt if r.font.size else "Default"
            b = "B" if r.font.bold else "Reg"
            run_info.append(f"[{fn}, {fs}pt, {b}]")
        pf = p.paragraph_format
        sb = pf.space_before.pt if pf.space_before else 0
        sa = pf.space_after.pt if pf.space_after else 0
        ls = pf.line_spacing if pf.line_spacing else "Default"
        print(f"{label} (P{idx:02d}): Align={p.alignment}, LineSp={ls}, SpB={sb}pt, SpA={sa}pt | Runs: {' '.join(run_info[:3])} | Text='{p.text[:40]}'")
