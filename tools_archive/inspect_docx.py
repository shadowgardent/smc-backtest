import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r"C:\Users\ONEST\Downloads\SMC_AUCC2026_Final.docx")
print("Total paragraphs:", len(doc.paragraphs))
print("Total tables:", len(doc.tables))

for i, p in enumerate(doc.paragraphs):
    has_drawing = any("drawing" in r._r.xml for r in p.runs)
    if has_drawing or "ภาพที่" in p.text or "ตารางที่" in p.text or p.text.startswith("1.") or p.text.startswith("2.") or p.text.startswith("3.") or p.text.startswith("4.") or p.text.startswith("5.") or p.text.startswith("6."):
        print(f"P{i:03d} (align={p.alignment}, drawing={has_drawing}): {p.text[:80]}")
