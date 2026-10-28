import docx, re, sys, subprocess, pymupdf
from docx.oxml import parse_xml

sys.stdout.reconfigure(encoding='utf-8')

backup_file = 'เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม_SMC_Revised_AUC2_Backup.docx'
doc = docx.Document(backup_file)

# 1. Update intro paragraph text (P39)
p39 = doc.paragraphs[39]
p39.text = p39.text.replace('ตามเงื่อนไขสมการ (1) และ (2):', 'ตามเงื่อนไขสมการ (1) และ (2) สำหรับทุกค่า i, j ∈ [1, N]:')

# Helper to remove the forall part from equation XML
def remove_forall(xml, target_sz):
    # Pattern from <m:r> with <m:t> </m:t> through </m:d>
    pattern = r'<m:r>.*?<m:t> </m:t>.*?</m:d>'
    xml_cleaned = re.sub(pattern, '', xml, flags=re.DOTALL)
    # update font sizes
    xml_resized = xml_cleaned.replace('w:val="12"', f'w:val="{target_sz}"')
    return xml_resized

for sz in [18, 20, 22]:
    d = docx.Document(backup_file)
    d.paragraphs[39].text = d.paragraphs[39].text.replace('ตามเงื่อนไขสมการ (1) และ (2):', 'ตามเงื่อนไขสมการ (1) และ (2) สำหรับทุกค่า i, j ∈ [1, N]:')
    
    # P40 (eq 1)
    p40_xml = remove_forall(d.paragraphs[40]._element.xml, sz)
    new_p40 = parse_xml(p40_xml.encode('utf-8'))
    d.paragraphs[40]._element.getparent().replace(d.paragraphs[40]._element, new_p40)
    
    # P41 (eq 1 label)
    p41_xml = d.paragraphs[41]._element.xml.replace('w:val="12"', f'w:val="{sz}"')
    new_p41 = parse_xml(p41_xml.encode('utf-8'))
    d.paragraphs[41]._element.getparent().replace(d.paragraphs[41]._element, new_p41)

    # P42 (eq 2)
    p42_xml = remove_forall(d.paragraphs[42]._element.xml, sz)
    new_p42 = parse_xml(p42_xml.encode('utf-8'))
    d.paragraphs[42]._element.getparent().replace(d.paragraphs[42]._element, new_p42)

    # P43 (eq 2 label)
    p43_xml = d.paragraphs[43]._element.xml.replace('w:val="12"', f'w:val="{sz}"')
    new_p43 = parse_xml(p43_xml.encode('utf-8'))
    d.paragraphs[43]._element.getparent().replace(d.paragraphs[43]._element, new_p43)

    # P52, P53 (eq 3)
    p52_xml = d.paragraphs[52]._element.xml.replace('w:val="12"', 'w:val="22"')
    new_p52 = parse_xml(p52_xml.encode('utf-8'))
    d.paragraphs[52]._element.getparent().replace(d.paragraphs[52]._element, new_p52)

    p53_xml = d.paragraphs[53]._element.xml.replace('w:val="12"', 'w:val="22"')
    new_p53 = parse_xml(p53_xml.encode('utf-8'))
    d.paragraphs[53]._element.getparent().replace(d.paragraphs[53]._element, new_p53)

    fn_docx = f'clean_sz_{sz}.docx'
    fn_pdf = f'clean_sz_{sz}.pdf'
    d.save(fn_docx)
    
    cmd = f'powershell -ExecutionPolicy Bypass -File e:\\Project\\smc-backtest-backend\\convert_docx_to_pdf.ps1 -docxPath "e:\\Project\\smc-backtest-backend\\{fn_docx}" -pdfPath "e:\\Project\\smc-backtest-backend\\{fn_pdf}"'
    subprocess.run(cmd, shell=True, check=True)
    
    doc_pdf = pymupdf.open(fn_pdf)
    pix = doc_pdf[2].get_pixmap(dpi=150)
    pix.save(f'clean_page3_sz_{sz}.png')
    print(f"Saved clean_page3_sz_{sz}.png")
