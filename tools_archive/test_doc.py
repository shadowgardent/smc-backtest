import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("e:/Project/smc-backtest-backend/SMC_AUCC2026_Final.docx")
for i, p in enumerate(doc.paragraphs):
    if "ภาพที่ 1" in p.text:
        print(f"P{i}: {p.text}")
        if i > 0:
            for r in doc.paragraphs[i-1].runs:
                if "drawing" in r._r.xml:
                    print("Found drawing in run before caption!")
