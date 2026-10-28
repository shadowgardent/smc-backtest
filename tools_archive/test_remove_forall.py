import docx, sys, subprocess, pymupdf
from docx.oxml import parse_xml
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

backup_file = 'เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม_SMC_Revised_AUC2_Backup.docx'

def clean_equation_element(p_elem, target_sz):
    # Parse with ET
    xml_str = p_elem.xml
    root = ET.fromstring(xml_str)
    ns = {
        'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    }
    
    # In row 0 (the first m:mr), find and remove the m:r or m:d that contain \u2003, ∀, or 1,N
    mr0 = root.find('.//m:mr', ns)
    if mr0 is not None:
        to_remove = []
        for child in mr0.findall('./m:e/*', ns):
            texts = [t.text for t in child.findall('.//m:t', ns) if t.text]
            t_combined = ''.join(texts)
            if any(k in t_combined for k in ['\u2003', '∀', '1,N']):
                to_remove.append(child)
        for c in to_remove:
            # find parent of c
            for parent in mr0.iter():
                if c in list(parent):
                    parent.remove(c)
                    break

    # update all w:sz to target_sz
    for sz in root.findall('.//w:sz', ns):
        sz.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', str(target_sz))
    for szCs in root.findall('.//w:szCs', ns):
        szCs.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', str(target_sz))

    new_xml = ET.tostring(root, encoding='utf-8').decode('utf-8')
    return parse_xml(new_xml.encode('utf-8'))

for sz in [16, 18, 20]:
    d = docx.Document(backup_file)
    d.paragraphs[39].text = d.paragraphs[39].text.replace(
        'ตามเงื่อนไขสมการ (1) และ (2):',
        'ตามเงื่อนไขสมการ (1) และ (2) สำหรับทุกค่า i, j ∈ [1, N]:'
    )
    
    new_p40 = clean_equation_element(d.paragraphs[40]._element, sz)
    d.paragraphs[40]._element.getparent().replace(d.paragraphs[40]._element, new_p40)

    new_p42 = clean_equation_element(d.paragraphs[42]._element, sz)
    d.paragraphs[42]._element.getparent().replace(d.paragraphs[42]._element, new_p42)

    # eq labels
    for p_idx in [41, 43]:
        p = d.paragraphs[p_idx]
        p_xml = p._element.xml.replace('w:val="12"', f'w:val="{sz}"')
        new_p = parse_xml(p_xml.encode('utf-8'))
        p._element.getparent().replace(p._element, new_p)

    # eq 3 (P52, 53)
    for p_idx in [52, 53]:
        p = d.paragraphs[p_idx]
        p_xml = p._element.xml.replace('w:val="12"', 'w:val="22"')
        new_p = parse_xml(p_xml.encode('utf-8'))
        p._element.getparent().replace(p._element, new_p)

    out_docx = f'clean_eq_{sz}.docx'
    out_pdf = f'clean_eq_{sz}.pdf'
    d.save(out_docx)
    
    cmd = f'powershell -ExecutionPolicy Bypass -File e:\\Project\\smc-backtest-backend\\convert_docx_to_pdf.ps1 -docxPath "e:\\Project\\smc-backtest-backend\\{out_docx}" -pdfPath "e:\\Project\\smc-backtest-backend\\{out_pdf}"'
    subprocess.run(cmd, shell=True, check=True)
    
    doc_pdf = pymupdf.open(out_pdf)
    pix = doc_pdf[2].get_pixmap(dpi=150)
    pix.save(f'clean_eq_page3_{sz}.png')
    print(f"Saved clean_eq_page3_{sz}.png")
