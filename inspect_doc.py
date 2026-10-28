import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

path = r"C:\Users\ONEST\Downloads\เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม Smart Money Concepts (SMC).docx"
doc = docx.Document(path)
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total sections: {len(doc.sections)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt:
        style_name = p.style.name if p.style else "NoStyle"
        print(f"P{i:03d} [{style_name}]: {txt}")
