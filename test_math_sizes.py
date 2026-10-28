import docx, re, sys, os, subprocess
from docx.shared import Inches, Pt
from docx.oxml import parse_xml

sys.stdout.reconfigure(encoding='utf-8')

backup_file = 'เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม_SMC_Revised_AUC2_Backup.docx'

for sz_val in [16, 18, 20]:
    doc = docx.Document(backup_file)
    for p_idx in [40, 41, 42, 43]:
        p = doc.paragraphs[p_idx]
        xml = p._element.xml
        new_xml = xml.replace('w:val="12"', f'w:val="{sz_val}"')
        new_element = parse_xml(new_xml.encode('utf-8'))
        p._element.getparent().replace(p._element, new_element)
    
    # Equation 3 can be 22 (11pt)
    for p_idx in [52, 53]:
        p = doc.paragraphs[p_idx]
        xml = p._element.xml
        new_xml = xml.replace('w:val="12"', 'w:val="22"')
        new_element = parse_xml(new_xml.encode('utf-8'))
        p._element.getparent().replace(p._element, new_element)
        
    test_docx = f'test_sz_{sz_val}.docx'
    test_pdf = f'test_sz_{sz_val}.pdf'
    doc.save(test_docx)
    
    cmd = f'powershell -ExecutionPolicy Bypass -File e:\\Project\\smc-backtest-backend\\convert_docx_to_pdf.ps1 -docxPath "e:\\Project\\smc-backtest-backend\\{test_docx}" -pdfPath "e:\\Project\\smc-backtest-backend\\{test_pdf}"'
    subprocess.run(cmd, shell=True, check=True)
    print(f"Generated {test_pdf}")
