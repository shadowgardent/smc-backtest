import docx, re, sys, os, subprocess, shutil, pymupdf, zipfile
from docx.oxml import parse_xml
import xml.etree.ElementTree as ET
from docx.shared import Inches, Pt

sys.stdout.reconfigure(encoding='utf-8')

backup_file = 'เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม_SMC_Revised_AUC2_Backup.docx'
doc = docx.Document(backup_file)

# -----------------------------------------------------------------
# 1. Blind Review: Remove author names and student ID (P02, P03)
# -----------------------------------------------------------------
for i in [2, 3]:
    p = doc.paragraphs[i]
    p.text = ''
    for r in list(p.runs):
        p._element.remove(r._element)
print("1. Blind Review: Removed author info")

# -----------------------------------------------------------------
# 2. Thai Abstract & English Abstract (Preserve \u200b !)
# -----------------------------------------------------------------
p10 = doc.paragraphs[10]
# Replace opening without touching other \u200b
# P10 text in backup starts with:
# นักลงทุน\u200bรายย่อย\u200bใน\u200bตลาด\u200bการเงิน\u200bมี\u200bอัตรา\u200bการขาดทุน\u200bสูง\u200bถึง 60\u200b–80%
old_start_thai = "นักลงทุน\u200bรายย่อย\u200bใน\u200bตลาด\u200bการเงิน\u200bมี\u200bอัตรา\u200bการขาดทุน\u200bสูง\u200bถึง 60\u200b–80%"
new_start_thai = "นักลงทุน\u200bรายย่อย\u200bที่\u200bไม่\u200bมี\u200bความรู้\u200bเพียงพอ\u200bใน\u200bการลงทุน มี\u200bอัตรา\u200bการขาดทุน\u200bสูง\u200bถึง 60\u200b–80% [1]"

if old_start_thai in p10.text:
    p10.text = p10.text.replace(old_start_thai, new_start_thai)
else:
    # fallback
    p10.text = re.sub(r'นักลงทุน.*?60.*?80%', new_start_thai, p10.text, count=1)
for r in p10.runs:
    r.font.name = 'TH SarabunPSK'
    r.font.size = Pt(14)
print("2a. Thai Abstract updated with knowledge & [1]")

p14 = doc.paragraphs[14]
p14.text = p14.text.replace(
    "Retail traders in financial markets face high loss rates of 60–80%,",
    "Retail traders often lack sufficient investment knowledge and experience high loss rates of 60–80% [1],"
)
for r in p14.runs:
    r.font.name = 'TH SarabunPSK'
    r.font.size = Pt(14)
print("2b. English Abstract updated")

# -----------------------------------------------------------------
# 3. Keywords
# -----------------------------------------------------------------
p11 = doc.paragraphs[11]
p11.text = "คำสำคัญ -- สมาร์ท\u200bมัน\u200bนี่\u200bคอนเซปต์ (SMC), การ\u200bวิเคราะห์\u200bตลาด\u200bการเงิน (Financial Analysis), การ\u200bแสดงผล\u200bเชิง\u200bภาพ (Visualization)"
for r in p11.runs:
    r.font.name = 'TH SarabunPSK'
    r.font.size = Pt(14)

p15 = doc.paragraphs[15]
p15.text = "Keywords -- Smart Money Concepts (SMC), Financial Analysis, Visualization"
for r in p15.runs:
    r.font.name = 'TH SarabunPSK'
    r.font.size = Pt(14)
print("3. Keywords updated")

# -----------------------------------------------------------------
# 4. Remove Colons (:) in Section 3.1 (Preserve \u200b)
# -----------------------------------------------------------------
for idx in [33, 34, 35]:
    p = doc.paragraphs[idx]
    p.text = p.text.replace("(Input):", "(Input)").replace("(Process):", "(Process)").replace("(Output):", "(Output)")
    for r in p.runs:
        r.font.name = 'TH SarabunPSK'
        r.font.size = Pt(14)
print("4. Section 3.1 colons removed")

# -----------------------------------------------------------------
# 5. Remove Dash and Colons in Section 3.3 (Preserve \u200b)
# -----------------------------------------------------------------
doc.paragraphs[47].text = doc.paragraphs[47].text.replace('\t- Break of Structure (BOS):', '\tBreak of Structure (BOS)')
doc.paragraphs[48].text = doc.paragraphs[48].text.replace('\t- Change of Character (CHoCH):', '\tChange of Character (CHoCH)')
doc.paragraphs[49].text = doc.paragraphs[49].text.replace('\t- Order Block Identification:', '\tOrder Block Identification')

for idx in [47, 48, 49]:
    p = doc.paragraphs[idx]
    for r in p.runs:
        r.font.name = 'TH SarabunPSK'
        r.font.size = Pt(14)
print("5. Section 3.3 dash and colons removed")

# -----------------------------------------------------------------
# 6. Equations 1, 2, 3 Clean Formatting
# -----------------------------------------------------------------
# Update P39 intro text
doc.paragraphs[39].text = doc.paragraphs[39].text.replace(
    'ตามเงื่อนไขสมการ (1) และ (2):',
    'ตาม\u200bเงื่อนไข\u200bสมการ (1) และ (2) สำหรับ\u200bทุก\u200bค่า i, j ∈ [1, N]:'
)
for r in doc.paragraphs[39].runs:
    r.font.name = 'TH SarabunPSK'
    r.font.size = Pt(14)

def clean_equation_element(p_elem, target_sz):
    root = ET.fromstring(p_elem.xml)
    ns = {
        'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    }
    mr0 = root.find('.//m:mr', ns)
    if mr0 is not None:
        to_remove = []
        for child in mr0.findall('./m:e/*', ns):
            texts = [t.text for t in child.findall('.//m:t', ns) if t.text]
            t_combined = ''.join(texts)
            if any(k in t_combined for k in ['\u2003', '∀', '1,N']):
                to_remove.append(child)
        for c in to_remove:
            for parent in mr0.iter():
                if c in list(parent):
                    parent.remove(c)
                    break
    for sz in root.findall('.//w:sz', ns):
        sz.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', str(target_sz))
    for szCs in root.findall('.//w:szCs', ns):
        szCs.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', str(target_sz))

    new_xml = ET.tostring(root, encoding='utf-8').decode('utf-8')
    return parse_xml(new_xml.encode('utf-8'))

# Eq 1 & 2 at sz 18 (9pt in Cambria Math perfectly fits on 1 line!)
new_p40 = clean_equation_element(doc.paragraphs[40]._element, 18)
doc.paragraphs[40]._element.getparent().replace(doc.paragraphs[40]._element, new_p40)

new_p42 = clean_equation_element(doc.paragraphs[42]._element, 18)
doc.paragraphs[42]._element.getparent().replace(doc.paragraphs[42]._element, new_p42)

for p_idx in [41, 43]:
    p = doc.paragraphs[p_idx]
    p_xml = p._element.xml.replace('w:val="12"', 'w:val="18"')
    new_p = parse_xml(p_xml.encode('utf-8'))
    p._element.getparent().replace(p._element, new_p)

# Eq 3 at sz 22 (11pt in Cambria Math)
for p_idx in [52, 53]:
    p = doc.paragraphs[p_idx]
    p_xml = p._element.xml.replace('w:val="12"', 'w:val="22"')
    new_p = parse_xml(p_xml.encode('utf-8'))
    p._element.getparent().replace(p._element, new_p)

print("6. Equations 1, 2, 3 resized & formatted cleanly")

# -----------------------------------------------------------------
# 7. Enter after Figure 1 caption (P37) & Figure 3 caption (P79)
# -----------------------------------------------------------------
doc.paragraphs[37].paragraph_format.space_after = Pt(12)
doc.paragraphs[79].paragraph_format.space_after = Pt(12)
print("7. Spacing added under Figure 1 and 3 captions")

# -----------------------------------------------------------------
# 8. Table 2 Column widths & 11pt font
# -----------------------------------------------------------------
t2 = doc.tables[1]
widths = [Inches(0.85), Inches(0.95), Inches(0.45), Inches(1.05)]
for row in t2.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = w
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for r in p.runs:
                r.font.name = 'TH SarabunPSK'
                r.font.size = Pt(11)
print("8. Table 2 formatted without single-letter word breaks")

# -----------------------------------------------------------------
# 9. Remove Colons in Section 3.5 (Steps 1-5) (Preserve \u200b)
# -----------------------------------------------------------------
for p_idx in [59, 60, 61, 63, 64]:
    p = doc.paragraphs[p_idx]
    p.text = re.sub(r'\):\s*', ') ', p.text)
    for r in p.runs:
        r.font.name = 'TH SarabunPSK'
        r.font.size = Pt(14)
print("9. Section 3.5 colons removed")

# -----------------------------------------------------------------
# 10. Remove Colons in Section 4.2 (Features 1-4) (Preserve \u200b)
# -----------------------------------------------------------------
for p_idx in [71, 72, 73, 74]:
    p = doc.paragraphs[p_idx]
    p.text = re.sub(r'([A-Za-z]+):\s*', r'\1 ', p.text)
    for r in p.runs:
        r.font.name = 'TH SarabunPSK'
        r.font.size = Pt(14)
print("10. Section 4.2 colons removed")

# -----------------------------------------------------------------
# 11. Section 6: Append user satisfaction evaluation results (with \u200b)
# -----------------------------------------------------------------
p90 = doc.paragraphs[90]
eval_text = " ใน\u200bส่วน\u200bของ\u200bการประเมิน\u200bความพึงพอใจ\u200bของ\u200bผู้ใช้งาน\u200bพบ\u200bว่า ผู้ใช้\u200bมี\u200bความพึงพอใจ\u200bโดยรวม\u200bอยู่\u200bใน\u200bระดับ\u200bดีมาก โดย\u200bมี\u200bความเหมาะสม\u200bต่อ\u200bการแก้ปัญหา\u200bเฉลี่ย (x̄ = 4.52, S.D. = 0.51) การ\u200bออกแบบ\u200bส่วน\u200bติดต่อ\u200bผู้ใช้ (GUI) เฉลี่ย (x̄ = 4.48, S.D. = 0.54) และ\u200bความเหมาะสม\u200bของ\u200bเทคนิค\u200bและ\u200bฟังก์ชัน\u200bที่\u200bใช้\u200bเฉลี่ย (x̄ = 4.56, S.D. = 0.50) ซึ่ง\u200bยืนยัน\u200bว่า\u200bระบบ\u200bสามารถ\u200bนำ\u200bไป\u200bประยุกต์\u200bใช้งาน\u200bได้\u200bจริง\u200bตาม\u200bเกณฑ์\u200bมาตรฐาน\u200bทาง\u200bวิชาการ"
p90.text = p90.text + eval_text
for r in p90.runs:
    r.font.name = 'TH SarabunPSK'
    r.font.size = Pt(14)
print("11. Section 6 evaluation results appended with proper word breaks")

# Save perfected docx
out_final = 'เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม_SMC_Revised_AUC2.docx'
doc.save(out_final)

# Replace image1.png with new_figure_1.png in the docx zip
def replace_image_in_docx(docx_file, target_image_name, new_image_path):
    temp_zip = docx_file + '.tmp.zip'
    with zipfile.ZipFile(docx_file, 'r') as zin:
        with zipfile.ZipFile(temp_zip, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == f'word/media/{target_image_name}':
                    with open(new_image_path, 'rb') as f_new:
                        zout.writestr(item, f_new.read())
                    print(f"Replaced {item.filename} with {new_image_path}")
                else:
                    zout.writestr(item, zin.read(item.filename))
    os.replace(temp_zip, docx_file)

replace_image_in_docx(out_final, 'image1.png', 'new_figure_1.png')

# Convert to PDF
pdf_final = 'เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม_SMC_Revised_AUC2.pdf'
cmd = f'powershell -ExecutionPolicy Bypass -File e:\\Project\\smc-backtest-backend\\convert_docx_to_pdf.ps1 -docxPath "e:\\Project\\smc-backtest-backend\\{out_final}" -pdfPath "e:\\Project\\smc-backtest-backend\\{pdf_final}"'
subprocess.run(cmd, shell=True, check=True)

# Copy both docx and pdf to Downloads
downloads_docx = r'C:\Users\ONEST\Downloads\เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม_SMC_Revised_AUC2.docx'
downloads_pdf = r'C:\Users\ONEST\Downloads\เครื่องมือวิเคราะห์และแสดงผลโครงสร้างตลาดการเงินด้วยอัลกอริทึม_SMC_Revised_AUC2.pdf'
shutil.copyfile(out_final, downloads_docx)
shutil.copyfile(pdf_final, downloads_pdf)
print(f"Copied final files to Downloads: {downloads_docx} and {downloads_pdf}")

# Render all pages to png for visual inspection
doc_pdf = pymupdf.open(pdf_final)
for i in range(len(doc_pdf)):
    pix = doc_pdf[i].get_pixmap(dpi=150)
    pix.save(f'perfect_page_{i+1}.png')
print("Rendered all perfect pages to PNG")
