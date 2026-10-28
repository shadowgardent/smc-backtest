import docx
from docx.shared import Cm

doc = docx.Document(r"C:\Users\ONEST\Downloads\SMC_AUCC2026_Final.docx")
p29 = doc.paragraphs[29]

# Replace image in P29 with SMC_Architecture_Flow_White.png with exact same width 9.38 cm
p29.text = "" # clear run
p29.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
run = p29.add_run()
run.add_picture(r"C:\Users\ONEST\Downloads\SMC_Architecture_Flow_White.png", width=Cm(9.38))

out_path = r"C:\Users\ONEST\Downloads\SMC_AUCC2026_Final_HD.docx"
doc.save(out_path)
print("Saved HD docx to:", out_path)
